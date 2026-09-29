"""A one-cell edit must diff to that one edit, not to cells nobody touched (issue #42).

`marinade diff` reported unedited cells as changed. Each mechanism is pinned
here against one small workbook, diffed against a copy with exactly one cell
changed:

- table candidates compared their members by binding_id, which is seeded with
  the workbook's sha256 -- so every multi-member table "changed" on any edit.
- a binding's label-scan context records the cached values of its neighbours,
  so editing one input re-hashed the bindings around it; that surfaced as
  BINDING_METADATA_CHANGED with no layer, i.e. as a workbook edit.

`Calc!B4:B5` look up `Data` by the key in `Calc!B1`; `Proj` is three parallel
formula columns, which is what the extractor groups into a table candidate.
"""

import sqlite3
from collections.abc import Callable
from pathlib import Path

import openpyxl
import pytest

import xl_marinade


def _workbook(path: Path, *, key: int = 3, rate: float = 0.05) -> Path:
    wb = openpyxl.Workbook()
    calc = wb.active
    calc.title = "Calc"
    calc["A1"], calc["B1"] = "Selected ID", key
    calc["A2"], calc["B2"] = "Rate", rate
    calc["A4"], calc["B4"] = "Value", "=INDEX(Data!$B$2:$B$11,MATCH($B$1,Data!$A$2:$A$11,0))"
    calc["A5"], calc["B5"] = "Weight", "=INDEX(Data!$C$2:$C$11,MATCH($B$1,Data!$A$2:$A$11,0))"
    calc["A6"], calc["B6"] = "Result", "=B4*B5*(1+B2)"

    data = wb.create_sheet("Data")
    for col, header in zip("ABC", ("ID", "Value", "Weight"), strict=True):
        data[f"{col}1"] = header
    for i in range(1, 11):
        data[f"A{i + 1}"], data[f"B{i + 1}"], data[f"C{i + 1}"] = i, i * 100, i / 10

    proj = wb.create_sheet("Proj")
    for col, header in zip("ABC", ("t", "Cashflow", "PV"), strict=True):
        proj[f"{col}1"] = header
    proj["A2"] = 0
    for r in range(2, 22):
        if r > 2:
            proj[f"A{r}"] = f"=A{r - 1}+1"
        proj[f"B{r}"] = f"=Calc!$B$6*(1+Calc!$B$2)^A{r}"
        proj[f"C{r}"] = f"=B{r}/1.03^A{r}"

    wb.save(path)
    return path


@pytest.fixture(scope="module")
def diff_against_base(tmp_path_factory: pytest.TempPathFactory) -> Callable[..., dict]:
    """diff(base, base-with-one-edit): the base is extracted once per module."""
    root = tmp_path_factory.mktemp("issue42")
    base_db = xl_marinade.extract(_workbook(root / "base.xlsx"), root / "base.db")
    cache: dict[str, dict] = {}

    def run(name: str, **edit: float) -> dict:
        if name not in cache:
            book = _workbook(root / f"{name}.xlsx", **edit)
            cache[name] = xl_marinade.diff(base_db, xl_marinade.extract(book, root / f"{name}.db"))
        return cache[name]

    return run


# Two single-cell edits: the lookup key, and an input that feeds no lookup.
EDITS = {"key": {"key": 5}, "rate": {"rate": 0.06}}


@pytest.mark.parametrize("edit", sorted(EDITS))
def test_unchanged_table_candidates_are_not_reported(
    diff_against_base: Callable[..., dict], edit: str
) -> None:
    result = diff_against_base(edit, **EDITS[edit])

    tables = [c for c in result["changes"] if c["type"].startswith("TABLE_CANDIDATE_")]
    assert tables == [], (
        "no table changed shape or membership, but the diff reported "
        f"{[(c['type'], c.get('candidate_id')) for c in tables]} -- members are being "
        "compared by workbook-scoped binding_id instead of binding position"
    )


def test_fixture_has_a_multi_member_table(tmp_path: Path) -> None:
    """Guard: without a multi-member candidate the test above proves nothing."""
    db = xl_marinade.extract(_workbook(tmp_path / "t.xlsx"), tmp_path / "t.db")
    (members,) = (
        sqlite3.connect(db)
        .execute(
            "SELECT MAX(n) FROM (SELECT COUNT(*) AS n FROM table_candidate_members GROUP BY candidate_id)"
        )
        .fetchone()
    )
    assert members and members > 1


@pytest.mark.parametrize("edit", sorted(EDITS))
def test_neighbour_metadata_churn_is_inference(
    diff_against_base: Callable[..., dict], edit: str
) -> None:
    result = diff_against_base(edit, **EDITS[edit])

    metadata = [c for c in result["changes"] if c["type"] == "BINDING_METADATA_CHANGED"]
    assert metadata, (
        "fixture no longer re-hashes a neighbouring binding's label context, so it "
        "cannot show that the churn is classified as inference"
    )
    assert {c.get("layer") for c in metadata} == {"ir_inference"}, [
        (c["address"], c.get("layer")) for c in metadata
    ]
